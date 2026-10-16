from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import parse_xml
from PIL import Image
Image.open('out/letterhead/_bg.png').convert('RGB').resize((1654,2339),Image.LANCZOS).save('bg_200.jpg',quality=92)
doc=Document(); sec=doc.sections[0]
sec.page_width,sec.page_height=Mm(210),Mm(297)
sec.top_margin,sec.bottom_margin,sec.right_margin,sec.left_margin=Mm(58),Mm(36),Mm(22),Mm(22)
sec.header_distance=Mm(0); sec.footer_distance=Mm(0)
# background picture anchored behind text in header
hp=sec.header.paragraphs[0]; run=hp.add_run(); run.add_picture('bg_200.jpg',width=Mm(210),height=Mm(297))
inline=run._r.xpath('.//wp:inline')[0]
graphic=inline.xpath('./a:graphic')[0]; extent=inline.xpath('./wp:extent')[0]; docpr=inline.xpath('./wp:docPr')[0]
from lxml import etree
anchor=parse_xml(f'''<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">
<wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH><wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>
<wp:extent cx="{extent.get('cx')}" cy="{extent.get('cy')}"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>
<wp:docPr id="{docpr.get('id')}" name="Letterhead"/><wp:cNvGraphicFramePr/></wp:anchor>''')
anchor.append(graphic); inline.getparent().replace(inline,anchor)
GREEN=RGBColor(0x0E,0x2E,0x2A)
def para(text,bold=False,align=WD_ALIGN_PARAGRAPH.RIGHT,size=12,color=None,space=6):
    p=doc.add_paragraph(); p.alignment={WD_ALIGN_PARAGRAPH.RIGHT:WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.LEFT:WD_ALIGN_PARAGRAPH.RIGHT}.get(align,align)
    pPr=p._p.get_or_add_pPr(); pPr.append(parse_xml('<w:bidi xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))
    p.paragraph_format.space_after=Pt(space); p.paragraph_format.line_spacing=1.5
    r=p.add_run(text); r.bold=bold; r.font.size=Pt(size); r.font.name='Tajawal'
    rPr=r._r.get_or_add_rPr(); rPr.rFonts.set(qn('w:cs'),'Tajawal'); rPr.append(parse_xml(f'<w:rtl xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))
    rPr.append(parse_xml(f'<w:szCs xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="{size*2}"/>'))
    if bold: rPr.append(parse_xml('<w:bCs xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))
    if color: r.font.color.rgb=color
    return p
L=WD_ALIGN_PARAGRAPH.LEFT
para('الرقم: ..................',size=10,align=L,space=0)
para('التاريخ: ..................',size=10,align=L,space=0)
para('المرفقات: ..................',size=10,align=L,space=12)
para('سعادة / ........................................   المحترم',bold=True,color=GREEN)
para('السلام عليكم ورحمة الله وبركاته، وبعد:')
para('اكتب نص الخطاب هنا...',align=WD_ALIGN_PARAGRAPH.JUSTIFY)
para('وتقبّلوا خالص التحية والتقدير،،،',align=WD_ALIGN_PARAGRAPH.CENTER,space=24)
para('فريق بودكاست نبض القصيم',bold=True,align=L,color=GREEN)
para('الاسم / التوقيع',align=L,size=10)
doc.save('out/letterhead/nabd-alqassim-letterhead.docx'); print('ok')
