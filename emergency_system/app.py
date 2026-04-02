"""
app.py - FastAPI Main Application
نظام قياس جاهزية الأقسام الحيوية اليومية
مستشفى الولادة والأطفال بالمدينة المنورة
"""
import os, io, base64, shutil, urllib.parse
from datetime import datetime, timedelta, date
from typing import Optional, List
from fastapi import (FastAPI, Depends, HTTPException, Request,
                     UploadFile, File, Form, Response)
from fastapi.responses import (HTMLResponse, JSONResponse,
                                FileResponse, StreamingResponse)
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

from database import (get_db, init_db, hash_password, get_setting, set_setting,
                      get_departments_template, get_bed_units_template,
                      get_report_by_date, get_report_by_id, get_reports_range,
                      DB_PATH)
from auth import (create_token, get_current_user, require_admin,
                  require_operator, authenticate_user)

BASE_DIR    = os.path.dirname(__file__)
STATIC_DIR  = os.path.join(BASE_DIR, "static")
LOGOS_DIR   = os.path.join(STATIC_DIR, "logos")
BACKUP_DIR  = os.path.join(STATIC_DIR, "backups")
FRONTEND_DIR= os.path.join(BASE_DIR, "frontend")

os.makedirs(LOGOS_DIR, exist_ok=True)
os.makedirs(BACKUP_DIR, exist_ok=True)

# ─── تهيئة قاعدة البيانات ─────────────────────────────────────────────────────
init_db()

app = FastAPI(title="نظام جاهزية الأقسام الحيوية", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"], allow_credentials=True,
    allow_methods=["*"], allow_headers=["*"],
)

app.mount("/static",   StaticFiles(directory=STATIC_DIR),   name="static")
app.mount("/frontend", StaticFiles(directory=FRONTEND_DIR), name="frontend")


# ═══════════════════════════════════════════════════════════════════════════════
# Pydantic Models
# ═══════════════════════════════════════════════════════════════════════════════

class LoginRequest(BaseModel):
    username: str
    password: str

class StaffDirectoryItem(BaseModel):
    name: str
    department: str
    phone: Optional[str] = ""
    role_title: Optional[str] = ""

class DepartmentTemplate(BaseModel):
    name: str
    requirement_label: Optional[str] = ""
    sort_order: Optional[int] = 0

class BedUnitTemplate(BaseModel):
    unit_name: str
    sort_order: Optional[int] = 0

class DailyStaffItem(BaseModel):
    department: str
    responsible_name: Optional[str] = ""
    phone: Optional[str] = ""
    requirement: Optional[str] = ""
    count_value: Optional[str] = ""
    readiness: Optional[str] = "جاهز"
    notes: Optional[str] = ""
    sort_order: Optional[int] = 0

class DailyBedItem(BaseModel):
    unit_name: str
    total_beds: Optional[int] = 0
    occupied: Optional[int] = 0
    sort_order: Optional[int] = 0

class DailyVentItem(BaseModel):
    on_vent: Optional[int] = 0
    ready_devices: Optional[int] = 0
    total_devices: Optional[int] = 0
    responsible_name: Optional[str] = ""
    responsible_phone: Optional[str] = ""

class SaveReportRequest(BaseModel):
    report_date: str
    shift: str
    notes: Optional[str] = ""
    staff: List[DailyStaffItem] = []
    beds: List[DailyBedItem] = []
    vents: Optional[DailyVentItem] = None

class UserCreate(BaseModel):
    username: str
    full_name: str
    password: str
    role: str

class UserUpdate(BaseModel):
    full_name: Optional[str] = None
    role: Optional[str] = None
    active: Optional[int] = None
    password: Optional[str] = None


# ═══════════════════════════════════════════════════════════════════════════════
# صفحات HTML
# ═══════════════════════════════════════════════════════════════════════════════

@app.get("/", response_class=HTMLResponse)
async def root():
    with open(os.path.join(FRONTEND_DIR, "index.html"), encoding="utf-8") as f:
        return f.read()

@app.get("/dashboard", response_class=HTMLResponse)
async def dashboard():
    with open(os.path.join(FRONTEND_DIR, "dashboard.html"), encoding="utf-8") as f:
        return f.read()


# ═══════════════════════════════════════════════════════════════════════════════
# Auth APIs
# ═══════════════════════════════════════════════════════════════════════════════

@app.post("/api/auth/login")
async def login(req: LoginRequest):
    user = authenticate_user(req.username, req.password)
    if not user:
        raise HTTPException(status_code=401,
                            detail="اسم المستخدم أو كلمة المرور غير صحيحة")
    token = create_token(user["id"], user["username"], user["role"])
    return {
        "token": token,
        "user": {
            "id": user["id"],
            "username": user["username"],
            "full_name": user["full_name"],
            "role": user["role"],
        }
    }

@app.get("/api/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return user


# ═══════════════════════════════════════════════════════════════════════════════
# Staff Directory APIs
# ═══════════════════════════════════════════════════════════════════════════════

@app.get("/api/staff")
async def list_staff(
    q: Optional[str] = None,
    department: Optional[str] = None,
    user: dict = Depends(get_current_user)
):
    conn = get_db()
    query = "SELECT * FROM staff_directory WHERE active=1"
    params = []
    if q:
        query += " AND (name LIKE ? OR department LIKE ? OR phone LIKE ?)"
        like = f"%{q}%"
        params += [like, like, like]
    if department:
        query += " AND department=?"
        params.append(department)
    query += " ORDER BY department, name"
    rows = conn.execute(query, params).fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.post("/api/staff")
async def create_staff(item: StaffDirectoryItem,
                       user: dict = Depends(require_operator)):
    conn = get_db()
    conn.execute(
        "INSERT INTO staff_directory (name, department, phone, role_title) VALUES (?,?,?,?)",
        (item.name, item.department, item.phone, item.role_title)
    )
    conn.commit()
    row_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
    row = conn.execute("SELECT * FROM staff_directory WHERE id=?", (row_id,)).fetchone()
    conn.close()
    return dict(row)

@app.put("/api/staff/{staff_id}")
async def update_staff(staff_id: int, item: StaffDirectoryItem,
                       user: dict = Depends(require_operator)):
    conn = get_db()
    conn.execute(
        "UPDATE staff_directory SET name=?, department=?, phone=?, role_title=? WHERE id=?",
        (item.name, item.department, item.phone, item.role_title, staff_id)
    )
    conn.commit()
    row = conn.execute("SELECT * FROM staff_directory WHERE id=?", (staff_id,)).fetchone()
    conn.close()
    if not row:
        raise HTTPException(404, "الكادر غير موجود")
    return dict(row)

@app.delete("/api/staff/{staff_id}")
async def delete_staff(staff_id: int, user: dict = Depends(require_operator)):
    conn = get_db()
    conn.execute("UPDATE staff_directory SET active=0 WHERE id=?", (staff_id,))
    conn.commit()
    conn.close()
    return {"message": "تم الحذف"}

@app.post("/api/staff/import")
async def import_staff(file: UploadFile = File(...),
                       user: dict = Depends(require_operator)):
    """استيراد كوادر من Excel"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise HTTPException(500, "openpyxl غير مثبت")
    content = await file.read()
    wb = load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.active
    headers = []
    inserted = updated = skipped = 0
    # كشف الأعمدة
    col_map = {}
    for row in ws.iter_rows(max_row=1, values_only=True):
        for idx, cell in enumerate(row):
            if cell:
                cell_lower = str(cell).strip().lower()
                if "الاسم" in str(cell) or "name" in cell_lower:
                    col_map["name"] = idx
                elif "قسم" in str(cell) or "dept" in cell_lower or "department" in cell_lower:
                    col_map["department"] = idx
                elif "جوال" in str(cell) or "هاتف" in str(cell) or "phone" in cell_lower or "تواصل" in str(cell):
                    col_map["phone"] = idx
                elif "مسمى" in str(cell) or "وظيف" in str(cell) or "title" in cell_lower or "role" in cell_lower:
                    col_map["role_title"] = idx
    if "name" not in col_map:
        raise HTTPException(400, "لم يتم العثور على عمود الاسم في الملف")
    conn = get_db()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        name = str(row[col_map.get("name", 0)] or "").strip()
        if not name:
            skipped += 1
            continue
        dept  = str(row[col_map.get("department", 1)] or "").strip() if "department" in col_map else ""
        phone = str(row[col_map.get("phone", 2)] or "").strip() if "phone" in col_map else ""
        title = str(row[col_map.get("role_title", 3)] or "").strip() if "role_title" in col_map else ""
        existing = conn.execute(
            "SELECT id FROM staff_directory WHERE name=? AND department=?", (name, dept)
        ).fetchone()
        if existing:
            conn.execute(
                "UPDATE staff_directory SET phone=?, role_title=?, active=1 WHERE id=?",
                (phone, title, existing["id"])
            )
            updated += 1
        else:
            conn.execute(
                "INSERT INTO staff_directory (name, department, phone, role_title) VALUES (?,?,?,?)",
                (name, dept, phone, title)
            )
            inserted += 1
    conn.commit()
    conn.close()
    return {"inserted": inserted, "updated": updated, "skipped": skipped}

@app.get("/api/staff/template/download")
async def download_staff_template(user: dict = Depends(get_current_user)):
    """تنزيل قالب Excel فارغ للاستيراد"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        raise HTTPException(500, "openpyxl غير مثبت")
    wb = Workbook()
    ws = wb.active
    ws.title = "قالب الكوادر"
    ws.sheet_view.rightToLeft = True
    headers = ["الاسم", "القسم", "رقم التواصل", "المسمى الوظيفي"]
    widths  = [30, 30, 20, 30]
    fill = PatternFill("solid", fgColor="1B4F8A")
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(bold=True, color="FFFFFF", name="Cairo")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   readingOrder=2)
        ws.column_dimensions[ws.cell(1, c).column_letter].width = w
    ws.row_dimensions[1].height = 25
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=staff_template.xlsx"}
    )


# ═══════════════════════════════════════════════════════════════════════════════
# Department Templates APIs
# ═══════════════════════════════════════════════════════════════════════════════

@app.get("/api/templates/departments")
async def get_dept_templates(user: dict = Depends(get_current_user)):
    return get_departments_template()

@app.post("/api/templates/departments")
async def add_dept_template(item: DepartmentTemplate,
                            user: dict = Depends(require_admin)):
    conn = get_db()
    conn.execute(
        "INSERT INTO departments_template (name, requirement_label, sort_order) VALUES (?,?,?)",
        (item.name, item.requirement_label, item.sort_order)
    )
    conn.commit()
    conn.close()
    return get_departments_template()

@app.put("/api/templates/departments/{dept_id}")
async def update_dept_template(dept_id: int, item: DepartmentTemplate,
                               user: dict = Depends(require_admin)):
    conn = get_db()
    conn.execute(
        "UPDATE departments_template SET name=?, requirement_label=?, sort_order=? WHERE id=?",
        (item.name, item.requirement_label, item.sort_order, dept_id)
    )
    conn.commit()
    conn.close()
    return get_departments_template()

@app.delete("/api/templates/departments/{dept_id}")
async def delete_dept_template(dept_id: int, user: dict = Depends(require_admin)):
    conn = get_db()
    conn.execute("UPDATE departments_template SET active=0 WHERE id=?", (dept_id,))
    conn.commit()
    conn.close()
    return {"message": "تم الحذف"}

@app.get("/api/templates/beds")
async def get_bed_templates(user: dict = Depends(get_current_user)):
    return get_bed_units_template()

@app.post("/api/templates/beds")
async def add_bed_template(item: BedUnitTemplate,
                           user: dict = Depends(require_admin)):
    conn = get_db()
    conn.execute(
        "INSERT INTO bed_units_template (unit_name, sort_order) VALUES (?,?)",
        (item.unit_name, item.sort_order)
    )
    conn.commit()
    conn.close()
    return get_bed_units_template()

@app.delete("/api/templates/beds/{unit_id}")
async def delete_bed_template(unit_id: int, user: dict = Depends(require_admin)):
    conn = get_db()
    conn.execute("UPDATE bed_units_template SET active=0 WHERE id=?", (unit_id,))
    conn.commit()
    conn.close()
    return {"message": "تم الحذف"}


# ═══════════════════════════════════════════════════════════════════════════════
# Daily Report APIs
# ═══════════════════════════════════════════════════════════════════════════════

@app.get("/api/reports/date/{report_date}")
async def get_report_date(report_date: str, user: dict = Depends(get_current_user)):
    report = get_report_by_date(report_date)
    if not report:
        raise HTTPException(404, "لا يوجد تقرير لهذا التاريخ")
    return report

@app.get("/api/reports/latest")
async def get_latest_report(user: dict = Depends(get_current_user)):
    conn = get_db()
    row = conn.execute(
        "SELECT report_date FROM daily_report ORDER BY report_date DESC LIMIT 1"
    ).fetchone()
    conn.close()
    if not row:
        raise HTTPException(404, "لا توجد تقارير")
    return get_report_by_date(row["report_date"])

@app.get("/api/reports/{report_id}")
async def get_report(report_id: int, user: dict = Depends(get_current_user)):
    report = get_report_by_id(report_id)
    if not report:
        raise HTTPException(404, "التقرير غير موجود")
    return report

@app.post("/api/reports/save")
async def save_report(req: SaveReportRequest,
                      user: dict = Depends(require_operator)):
    conn = get_db()
    try:
        # التحقق من وجود تقرير مسبق
        existing = conn.execute(
            "SELECT id FROM daily_report WHERE report_date=?", (req.report_date,)
        ).fetchone()

        if existing:
            report_id = existing["id"]
            conn.execute(
                "UPDATE daily_report SET shift=?, updated_by=?, notes=?, "
                "updated_at=datetime('now','localtime') WHERE id=?",
                (req.shift, user["id"], req.notes, report_id)
            )
            conn.execute("DELETE FROM daily_staff WHERE report_id=?", (report_id,))
            conn.execute("DELETE FROM daily_beds  WHERE report_id=?", (report_id,))
            conn.execute("DELETE FROM daily_vents WHERE report_id=?", (report_id,))
        else:
            conn.execute(
                "INSERT INTO daily_report (report_date, shift, created_by, updated_by, notes) "
                "VALUES (?,?,?,?,?)",
                (req.report_date, req.shift, user["id"], user["id"], req.notes)
            )
            report_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]

        # حفظ كوادر الشفت
        for s in req.staff:
            conn.execute("""
                INSERT INTO daily_staff
                (report_id, department, responsible_name, phone, requirement,
                 count_value, readiness, notes, sort_order)
                VALUES (?,?,?,?,?,?,?,?,?)
            """, (report_id, s.department, s.responsible_name, s.phone,
                  s.requirement, s.count_value, s.readiness, s.notes, s.sort_order))

        # حفظ الأسرة
        for b in req.beds:
            conn.execute("""
                INSERT INTO daily_beds
                (report_id, unit_name, total_beds, occupied, sort_order)
                VALUES (?,?,?,?,?)
            """, (report_id, b.unit_name, b.total_beds, b.occupied, b.sort_order))

        # حفظ أجهزة التنفس
        if req.vents:
            conn.execute("""
                INSERT INTO daily_vents
                (report_id, on_vent, ready_devices, total_devices,
                 responsible_name, responsible_phone)
                VALUES (?,?,?,?,?,?)
            """, (report_id, req.vents.on_vent, req.vents.ready_devices,
                  req.vents.total_devices, req.vents.responsible_name,
                  req.vents.responsible_phone))

        conn.commit()
        return {"message": "تم حفظ التقرير بنجاح", "report_id": report_id,
                "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S")}
    except Exception as e:
        conn.rollback()
        raise HTTPException(500, f"خطأ في حفظ التقرير: {str(e)}")
    finally:
        conn.close()

@app.get("/api/reports/archive/list")
async def list_archive(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    shift: Optional[str] = None,
    page: int = 1,
    per_page: int = 20,
    user: dict = Depends(get_current_user)
):
    conn = get_db()
    query  = "SELECT dr.*, u.full_name as creator_name FROM daily_report dr "
    query += "LEFT JOIN users u ON dr.created_by=u.id WHERE 1=1"
    params = []
    if start_date:
        query += " AND dr.report_date >= ?"; params.append(start_date)
    if end_date:
        query += " AND dr.report_date <= ?"; params.append(end_date)
    if shift:
        query += " AND dr.shift=?"; params.append(shift)
    query += " ORDER BY dr.report_date DESC"
    total = conn.execute(
        "SELECT COUNT(*) FROM (" + query + ")", params
    ).fetchone()[0]
    query += f" LIMIT {per_page} OFFSET {(page-1)*per_page}"
    rows = conn.execute(query, params).fetchall()
    conn.close()
    return {"total": total, "page": page, "per_page": per_page,
            "items": [dict(r) for r in rows]}

@app.delete("/api/reports/{report_id}")
async def delete_report(report_id: int, user: dict = Depends(require_admin)):
    conn = get_db()
    conn.execute("DELETE FROM daily_report WHERE id=?", (report_id,))
    conn.commit()
    conn.close()
    return {"message": "تم حذف التقرير"}


# ═══════════════════════════════════════════════════════════════════════════════
# Analytics APIs
# ═══════════════════════════════════════════════════════════════════════════════

@app.get("/api/analytics/summary")
async def analytics_summary(
    period: str = "weekly",
    ref_date: Optional[str] = None,
    user: dict = Depends(get_current_user)
):
    """ملخص إحصائي للفترة المحددة"""
    try:
        base = datetime.strptime(ref_date, "%Y-%m-%d").date() if ref_date \
               else date.today()
    except Exception:
        base = date.today()

    if period == "weekly":
        start = base - timedelta(days=6)
    elif period == "monthly":
        start = base.replace(day=1)
    else:
        start = base

    reports = get_reports_range(start.strftime("%Y-%m-%d"), base.strftime("%Y-%m-%d"))
    if not reports:
        return {"period": period, "start": str(start), "end": str(base),
                "reports_count": 0, "daily_stats": []}

    daily_stats = []
    dept_issues  = {}

    for r in reports:
        # إشغال الأسرة
        beds = r.get("beds", [])
        total_b = sum(b.get("total_beds", 0) or 0 for b in beds)
        occ_b   = sum(b.get("occupied", 0) or 0 for b in beds)
        occ_pct = round((occ_b/total_b*100) if total_b > 0 else 0, 1)

        # جاهزية Vent
        v = r.get("vents", {})
        total_v = v.get("total_devices", 0) or 0
        ready_v = v.get("ready_devices", 0) or 0
        vent_pct = round((ready_v/total_v*100) if total_v > 0 else 0, 1)

        # مؤشرات الكوادر
        staff     = r.get("staff", [])
        not_resp  = sum(1 for s in staff if s.get("readiness") == "لم يتم التجاوب")
        partial   = sum(1 for s in staff if s.get("readiness") == "تحت الاستجابة")

        for s in staff:
            if s.get("readiness") != "جاهز":
                d = s.get("department", "غير محدد")
                dept_issues[d] = dept_issues.get(d, 0) + 1

        daily_stats.append({
            "date":         r["report_date"],
            "shift":        r.get("shift",""),
            "occ_pct":      occ_pct,
            "vent_pct":     vent_pct,
            "not_responded":not_resp,
            "partial":      partial,
            "over_90":      occ_pct >= 90,
        })

    avg_occ  = round(sum(d["occ_pct"]  for d in daily_stats) / len(daily_stats), 1)
    avg_vent = round(sum(d["vent_pct"] for d in daily_stats) / len(daily_stats), 1)
    days_over_90 = sum(1 for d in daily_stats if d["over_90"])
    dept_issues_sorted = sorted(dept_issues.items(), key=lambda x: x[1], reverse=True)

    return {
        "period":          period,
        "start":           str(start),
        "end":             str(base),
        "reports_count":   len(reports),
        "avg_occ_pct":     avg_occ,
        "avg_vent_pct":    avg_vent,
        "days_over_90":    days_over_90,
        "dept_issues":     [{"dept": k, "count": v} for k, v in dept_issues_sorted],
        "daily_stats":     daily_stats,
    }


# ═══════════════════════════════════════════════════════════════════════════════
# Export APIs
# ═══════════════════════════════════════════════════════════════════════════════

def _get_logos_b64():
    """استرجاع الشعارات بصيغة base64"""
    def load_logo(name):
        path = os.path.join(LOGOS_DIR, name)
        if os.path.exists(path):
            with open(path, "rb") as f:
                return base64.b64encode(f.read()).decode()
        return ""
    return load_logo("hospital.png"), load_logo("emergency.png")

@app.get("/api/export/excel/{report_id}")
async def export_excel_report(report_id: int, user: dict = Depends(get_current_user)):
    report = get_report_by_id(report_id)
    if not report:
        raise HTTPException(404, "التقرير غير موجود")
    from export_utils import export_excel
    try:
        data = export_excel(report, get_setting("hospital_name", "مستشفى الولادة والأطفال"))
    except RuntimeError as e:
        raise HTTPException(500, str(e))
    filename = f"مؤشر_التحضير_{report['report_date']}.xlsx"
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename*=UTF-8''{}".format(urllib.parse.quote(filename))}
    )

@app.get("/api/export/pdf/{report_id}")
async def export_pdf_report(report_id: int, user: dict = Depends(get_current_user)):
    report = get_report_by_id(report_id)
    if not report:
        raise HTTPException(404, "التقرير غير موجود")
    from export_utils import export_pdf, build_report_html
    logo_b64, emg_logo_b64 = _get_logos_b64()
    hospital_name     = get_setting("hospital_name", "مستشفى الولادة والأطفال")
    hospital_subtitle = get_setting("hospital_subtitle", "إدارة التأهب للطوارئ والكوارث")
    try:
        data = export_pdf(report, hospital_name, hospital_subtitle,
                          logo_b64, emg_logo_b64)
        filename = f"تقرير_جاهزية_{report['report_date']}.pdf"
        return StreamingResponse(
            io.BytesIO(data),
            media_type="application/pdf",
            headers={"Content-Disposition": "attachment; filename*=UTF-8''{}".format(urllib.parse.quote(filename))}
        )
    except RuntimeError:
        # fallback: إرجاع HTML للطباعة
        html = build_report_html(report, hospital_name, hospital_subtitle,
                                 logo_b64, emg_logo_b64)
        return HTMLResponse(content=html)

@app.get("/api/export/html/{report_id}")
async def export_html_report(report_id: int, user: dict = Depends(get_current_user)):
    """إرجاع HTML للطباعة من المتصفح"""
    report = get_report_by_id(report_id)
    if not report:
        raise HTTPException(404, "التقرير غير موجود")
    from export_utils import build_report_html
    logo_b64, emg_logo_b64 = _get_logos_b64()
    html = build_report_html(
        report,
        get_setting("hospital_name", "مستشفى الولادة والأطفال"),
        get_setting("hospital_subtitle", "إدارة التأهب للطوارئ والكوارث"),
        logo_b64, emg_logo_b64
    )
    return HTMLResponse(content=html)


# ═══════════════════════════════════════════════════════════════════════════════
# Users Management APIs
# ═══════════════════════════════════════════════════════════════════════════════

@app.get("/api/users")
async def list_users(user: dict = Depends(require_admin)):
    conn = get_db()
    rows = conn.execute(
        "SELECT id, username, full_name, role, active, created_at FROM users ORDER BY id"
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.post("/api/users")
async def create_user(req: UserCreate, user: dict = Depends(require_admin)):
    if req.role not in ("admin", "operator", "viewer"):
        raise HTTPException(400, "دور غير صالح")
    conn = get_db()
    existing = conn.execute("SELECT id FROM users WHERE username=?", (req.username,)).fetchone()
    if existing:
        conn.close()
        raise HTTPException(400, "اسم المستخدم مستخدم مسبقاً")
    conn.execute(
        "INSERT INTO users (username, full_name, password_hash, role) VALUES (?,?,?,?)",
        (req.username, req.full_name, hash_password(req.password), req.role)
    )
    conn.commit()
    conn.close()
    return {"message": "تم إنشاء المستخدم بنجاح"}

@app.put("/api/users/{user_id}")
async def update_user(user_id: int, req: UserUpdate,
                      current_user: dict = Depends(require_admin)):
    conn = get_db()
    if req.full_name is not None:
        conn.execute("UPDATE users SET full_name=? WHERE id=?", (req.full_name, user_id))
    if req.role is not None:
        if req.role not in ("admin", "operator", "viewer"):
            conn.close()
            raise HTTPException(400, "دور غير صالح")
        conn.execute("UPDATE users SET role=? WHERE id=?", (req.role, user_id))
    if req.active is not None:
        conn.execute("UPDATE users SET active=? WHERE id=?", (req.active, user_id))
    if req.password:
        conn.execute("UPDATE users SET password_hash=? WHERE id=?",
                     (hash_password(req.password), user_id))
    conn.commit()
    conn.close()
    return {"message": "تم تحديث المستخدم"}


# ═══════════════════════════════════════════════════════════════════════════════
# Settings APIs
# ═══════════════════════════════════════════════════════════════════════════════

@app.get("/api/settings")
async def get_all_settings(user: dict = Depends(get_current_user)):
    conn = get_db()
    rows = conn.execute("SELECT key, value FROM settings").fetchall()
    conn.close()
    result = {r["key"]: r["value"] for r in rows}
    # إضافة URLs الشعارات
    result["hospital_logo_url"]  = (
        "/static/logos/hospital.png"
        if os.path.exists(os.path.join(LOGOS_DIR, "hospital.png")) else ""
    )
    result["emergency_logo_url"] = (
        "/static/logos/emergency.png"
        if os.path.exists(os.path.join(LOGOS_DIR, "emergency.png")) else ""
    )
    return result

@app.post("/api/settings")
async def update_settings(request: Request, user: dict = Depends(require_admin)):
    data = await request.json()
    for key, value in data.items():
        set_setting(key, str(value))
    return {"message": "تم حفظ الإعدادات"}

@app.post("/api/settings/logo/{logo_type}")
async def upload_logo(logo_type: str, file: UploadFile = File(...),
                      user: dict = Depends(require_admin)):
    if logo_type not in ("hospital", "emergency"):
        raise HTTPException(400, "نوع الشعار غير صحيح")
    allowed = {".png", ".jpg", ".jpeg", ".svg", ".webp"}
    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in allowed:
        raise HTTPException(400, "صيغة الملف غير مدعومة (PNG/JPG/SVG فقط)")
    # حفظ بصيغة PNG دائماً في الاسم
    save_path = os.path.join(LOGOS_DIR, f"{logo_type}.png")
    content = await file.read()
    with open(save_path, "wb") as f:
        f.write(content)
    return {"message": "تم رفع الشعار", "url": f"/static/logos/{logo_type}.png"}


# ═══════════════════════════════════════════════════════════════════════════════
# Backup & Restore APIs
# ═══════════════════════════════════════════════════════════════════════════════

@app.get("/api/backup/download")
async def download_backup(user: dict = Depends(require_admin)):
    if not os.path.exists(DB_PATH):
        raise HTTPException(404, "ملف قاعدة البيانات غير موجود")
    timestamp  = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_name= f"emergency_backup_{timestamp}.db"
    return FileResponse(DB_PATH, filename=backup_name,
                        media_type="application/octet-stream")

@app.post("/api/backup/restore")
async def restore_backup(file: UploadFile = File(...),
                         user: dict = Depends(require_admin)):
    if not file.filename.endswith(".db"):
        raise HTTPException(400, "يجب رفع ملف بصيغة .db")
    content = await file.read()
    # نسخة احتياطية قبل الاستبدال
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = os.path.join(BACKUP_DIR, f"pre_restore_{ts}.db")
    if os.path.exists(DB_PATH):
        shutil.copy2(DB_PATH, backup_path)
    with open(DB_PATH, "wb") as f:
        f.write(content)
    return {"message": "تم استعادة قاعدة البيانات بنجاح"}


# ═══════════════════════════════════════════════════════════════════════════════
# Run
# ═══════════════════════════════════════════════════════════════════════════════

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("app:app", host="127.0.0.1", port=8000, reload=False)
